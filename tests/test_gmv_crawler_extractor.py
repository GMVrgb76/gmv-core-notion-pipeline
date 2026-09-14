"""Crawler preplan step 10: Extractors (spec v0.2 §7).

Cross-checks STATUS_VALUES against the real EvidenceError codes
gmv_evidence_pipeline._extract()/_paddleocr_extract() can actually raise
(read from their real source via inspect, not a hand-typed duplicate list)
-- same discipline steps 4/6/7/8 already used for content_hash/entity_type/
predicate/gmv_id governance.

Deliberately does not re-test every extraction-format edge case already
covered by tests/test_gmv_evidence_pipeline.py (e.g. every PaddleOCR
subprocess failure mode, every soffice failure mode) -- gmv_crawler_extractor
wraps _extract() by direct import, so those cases are exercised by that
existing suite already; re-asserting them here would be a duplicated test,
not a real cross-check. This file tests only what gmv_crawler_extractor.py
adds on top: ExtractionDocument's own invariants, the source_hash staleness
gate, and status-code mapping.
"""
from __future__ import annotations

import inspect
import re
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

ROOT = Path(__file__).parents[1]
sys.path.insert(0, str(ROOT / "10_API"))
sys.path.insert(0, str(ROOT))

import gmv_crawler_extractor as ce  # noqa: E402
import gmv_evidence_pipeline as evidence  # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures"
GOOD_HASH = "sha256:" + "a" * 64


def _hash_of(path: Path) -> str:
    return f"sha256:{evidence.sha256_file(path)}"


# --- cross-check: STATUS_VALUES covers every real EvidenceError code
# _extract()/_paddleocr_extract() can raise ---


def test_status_values_cover_every_real_extract_error_code() -> None:
    source = inspect.getsource(evidence._extract) + inspect.getsource(evidence._paddleocr_extract)
    codes = set(re.findall(r'raise EvidenceError\(\s*"([A-Z_]+)"', source))
    assert codes, "regex found nothing -- _extract()'s real source changed shape, update the regex"
    assert codes <= ce.STATUS_VALUES, (
        f"_extract() can raise {codes - ce.STATUS_VALUES} which extract_document() would "
        "silently fold into EXTRACTION_FAILED -- add the real code to STATUS_VALUES instead"
    )


# --- ExtractionDocument invariants ---


def test_extraction_document_rejects_empty_source_id() -> None:
    with pytest.raises(ValueError, match="source_id"):
        ce.ExtractionDocument(source_id="", source_hash=GOOD_HASH, status="EXTRACTION_FAILED")


def test_extraction_document_rejects_malformed_source_hash() -> None:
    with pytest.raises(ValueError):
        ce.ExtractionDocument(source_id="SRC-1", source_hash="not-a-hash", status="EXTRACTION_FAILED")


def test_extraction_document_rejects_unknown_status() -> None:
    with pytest.raises(ValueError, match="status"):
        ce.ExtractionDocument(source_id="SRC-1", source_hash=GOOD_HASH, status="NOT_A_REAL_STATUS")


def test_extraction_document_success_requires_text() -> None:
    with pytest.raises(ValueError, match="text"):
        ce.ExtractionDocument(source_id="SRC-1", source_hash=GOOD_HASH, status="SUCCESS", extractor="text")


def test_extraction_document_success_requires_extractor() -> None:
    with pytest.raises(ValueError, match="extractor"):
        ce.ExtractionDocument(source_id="SRC-1", source_hash=GOOD_HASH, status="SUCCESS", text="hello")


def test_extraction_document_non_success_rejects_stray_text() -> None:
    with pytest.raises(ValueError, match="text"):
        ce.ExtractionDocument(source_id="SRC-1", source_hash=GOOD_HASH, status="EXTRACTION_FAILED", text="hello")


def test_extraction_document_non_success_rejects_stray_extractor() -> None:
    with pytest.raises(ValueError, match="extractor"):
        ce.ExtractionDocument(source_id="SRC-1", source_hash=GOOD_HASH, status="UNSUPPORTED_FORMAT", extractor="text")


def test_extraction_document_accepts_well_formed_success() -> None:
    doc = ce.ExtractionDocument(source_id="SRC-1", source_hash=GOOD_HASH, status="SUCCESS",
                                 extractor="text", text="hello")
    assert doc.language is None
    assert doc.structural_units == ()
    assert doc.warnings == ()


# --- extract_document(): staleness gate ---


def test_extract_document_reports_stale_hash_when_file_does_not_match(tmp_path: Path) -> None:
    path = tmp_path / "a.txt"; path.write_text("Federico Garibaldi", encoding="utf-8")
    doc = ce.extract_document(path, source_id="SRC-1", source_hash=GOOD_HASH)
    assert doc.status == "EXTRACTION_ABORTED_STALE_HASH"
    assert doc.text == "" and doc.extractor == ""


def test_extract_document_reports_stale_hash_when_file_missing(tmp_path: Path) -> None:
    doc = ce.extract_document(tmp_path / "missing.txt", source_id="SRC-1", source_hash=GOOD_HASH)
    assert doc.status == "EXTRACTION_ABORTED_STALE_HASH"


def test_extract_document_rejects_malformed_caller_hash(tmp_path: Path) -> None:
    path = tmp_path / "a.txt"; path.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        ce.extract_document(path, source_id="SRC-1", source_hash="not-a-hash")


# --- extract_document(): real delegation to _extract() (reuse, not reimplementation) ---


def test_extract_document_txt_success(tmp_path: Path) -> None:
    path = tmp_path / "note.txt"; path.write_text("Federico Garibaldi", encoding="utf-8")
    doc = ce.extract_document(path, source_id="SRC-1", source_hash=_hash_of(path))
    assert doc.status == "SUCCESS"
    assert doc.extractor == "text"
    assert doc.text == "Federico Garibaldi"
    assert doc.source_id == "SRC-1"
    assert doc.extractor_version == ce.EXTRACTOR_VERSION


def test_extract_document_md_success(tmp_path: Path) -> None:
    path = tmp_path / "bio.md"; path.write_text("# Federico Garibaldi", encoding="utf-8")
    doc = ce.extract_document(path, source_id="SRC-1", source_hash=_hash_of(path))
    assert doc.status == "SUCCESS"
    assert doc.extractor == "text"


def test_extract_document_pdf_native_text_layer(tmp_path: Path) -> None:
    path = tmp_path / "bio.pdf"; path.write_bytes((FIXTURES / "text_sample.pdf").read_bytes())
    doc = ce.extract_document(path, source_id="SRC-1", source_hash=_hash_of(path))
    assert doc.status == "SUCCESS"
    assert doc.extractor == "pdf_text"
    assert doc.text


def test_extract_document_pdf_corrupt_fails_explicitly(tmp_path: Path) -> None:
    path = tmp_path / "bad.pdf"; path.write_bytes((FIXTURES / "corrupt.pdf").read_bytes())
    doc = ce.extract_document(path, source_id="SRC-1", source_hash=_hash_of(path))
    assert doc.status == "EXTRACTION_FAILED"
    assert doc.text == "" and doc.extractor == ""


def test_extract_document_unsupported_format(tmp_path: Path) -> None:
    path = tmp_path / "archive.zip"; path.write_bytes(b"PK\x03\x04")
    doc = ce.extract_document(path, source_id="SRC-1", source_hash=_hash_of(path))
    assert doc.status == "UNSUPPORTED_FORMAT"


def test_extract_document_docx_success(tmp_path: Path) -> None:
    from docx import Document
    document = Document()
    document.add_paragraph("Federico Garibaldi")
    path = tmp_path / "bio.docx"
    document.save(path)
    doc = ce.extract_document(path, source_id="SRC-1", source_hash=_hash_of(path))
    assert doc.status == "SUCCESS"
    assert doc.extractor == "docx_text"
    assert "Federico Garibaldi" in doc.text


def test_extract_document_pdf_ocr_required_via_real_paddleocr_path(monkeypatch, tmp_path: Path) -> None:
    """Mirrors test_gmv_evidence_pipeline.py's own OCR_REQUIRED coverage
    (fake venv present, PaddleOCR subprocess returns blank text) -- proves
    extract_document() surfaces the same status _extract() actually
    produces, not a re-derivation of PaddleOCR's own logic."""
    fake_venv = tmp_path / "fake_venv_python"; fake_venv.write_text("")
    monkeypatch.setattr(evidence, "PADDLEOCR_VENV_PYTHON", fake_venv)
    monkeypatch.setattr(evidence.subprocess, "run", lambda *a, **k: SimpleNamespace(stdout="   \n", stderr=""))
    path = tmp_path / "scan.pdf"; path.write_bytes((FIXTURES / "scanned_blank.pdf").read_bytes())
    doc = ce.extract_document(path, source_id="SRC-1", source_hash=_hash_of(path))
    assert doc.status == "OCR_REQUIRED"


def test_extract_document_non_utf8_text_fails_explicitly_instead_of_crashing(tmp_path: Path) -> None:
    """Regression guard: _extract()'s plain-text branch reads with
    errors="strict" and no try/except of its own -- a non-UTF-8 .txt file
    raises a bare UnicodeDecodeError straight out of _extract(). Without
    catching (OSError, UnicodeError) too (not just EvidenceError),
    extract_document() would crash instead of reporting EXTRACTION_FAILED,
    unlike gmv_evidence_pipeline.extract() itself, which already guards
    against exactly this."""
    path = tmp_path / "latin1.txt"; path.write_bytes(b"caf\xe9")
    doc = ce.extract_document(path, source_id="SRC-1", source_hash=_hash_of(path))
    assert doc.status == "EXTRACTION_FAILED"
    assert doc.error_detail == "UnicodeDecodeError"
    assert doc.text == "" and doc.extractor == ""


def test_extract_document_never_raises_evidence_error(tmp_path: Path) -> None:
    """Every format-extraction failure must surface as a status, not an
    exception -- a caller iterating many sources must not need a
    per-file try/except (see extract_document()'s own docstring)."""
    path = tmp_path / "empty.pdf"; path.write_bytes((FIXTURES / "empty.pdf").read_bytes())
    doc = ce.extract_document(path, source_id="SRC-1", source_hash=_hash_of(path))
    assert doc.status in ce.STATUS_VALUES
